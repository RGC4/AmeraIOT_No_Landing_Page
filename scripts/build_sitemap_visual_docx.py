import math
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

IMG = "screenshots/sitemap/"

DARK = RGBColor(0x1F, 0x2A, 0x33)
GREY = RGBColor(0x66, 0x66, 0x66)
PRIMARY = RGBColor(0x2F, 0x6F, 0x9F)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def cell_border(cell, color="C9D6E0", sz="8"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def no_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tcPr.append(borders)


def fill_cell(cell, label, path, img_file, img_w):
    shade(cell, "FFFFFF")
    cell_border(cell)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(IMG + img_file, width=Inches(img_w))
    cap = cell.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(4)
    r1 = cap.add_run(label)
    r1.bold = True
    r1.font.size = Pt(9)
    r1.font.color.rgb = DARK
    r2 = cap.add_run("\n" + path)
    r2.font.size = Pt(7.5)
    r2.font.color.rgb = GREY


def add_grid(doc, items, cols, img_w):
    rows = math.ceil(len(items) / cols)
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, (label, path, img_file) in enumerate(items):
        r, c = divmod(i, cols)
        fill_cell(table.cell(r, c), label, path, img_file, img_w)
    # blank out unused trailing cells
    for j in range(len(items), rows * cols):
        r, c = divmod(j, cols)
        cell = table.cell(r, c)
        cell.text = ""
        no_border(cell)


def section_heading(doc, text, space_before=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = PRIMARY
    # bottom rule
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "C9D6E0")
    pbdr.append(bottom)
    pPr.append(pbdr)


# --- Document --------------------------------------------------------------
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(11), Inches(8.5)
sec.left_margin = sec.right_margin = Inches(0.4)
sec.top_margin = sec.bottom_margin = Inches(0.4)

# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = title.add_run("AMERA Website \u2014 Visual Site Map")
tr.bold = True
tr.font.size = Pt(20)
tr.font.color.rgb = DARK
title.paragraph_format.space_after = Pt(2)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sub.add_run("Page thumbnails arranged by site hierarchy")
sr.italic = True
sr.font.size = Pt(10)
sr.font.color.rgb = GREY
sub.paragraph_format.space_after = Pt(6)

# Home (centered single)
section_heading(doc, "Home  \u2014  /", space_before=4)
home_tbl = doc.add_table(rows=1, cols=1)
home_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
fill_cell(home_tbl.cell(0, 0), "Home", "/", "home.jpg", 5.0)

# Main navigation (6)
section_heading(doc, "Main Navigation")
add_grid(doc, [
    ("Platform", "/products", "products.jpg"),
    ("Industry Use Cases", "/industries", "industries.jpg"),
    ("Resources", "/resources", "resources.jpg"),
    ("News", "/news", "news.jpg"),
    ("Company", "/company", "company.jpg"),
    ("Contact Us", "/contact", "contact.jpg"),
], cols=3, img_w=3.05)

# Platform children
doc.add_page_break()
section_heading(doc, "Platform  \u2014  /products", space_before=4)
add_grid(doc, [
    ("AmeraKey", "/products/amerakey", "amerakey.jpg"),
    ("AmeraSecrets", "/products/amerasecrets", "amerasecrets.jpg"),
], cols=2, img_w=4.6)

# Company children
section_heading(doc, "Company  \u2014  /company")
add_grid(doc, [
    ("Vision & Mission", "/company/vision-and-mission", "vm.jpg"),
    ("Executive Team", "/company/executive-team", "exec.jpg"),
    ("Board of Advisors", "/company/board-of-advisors", "board.jpg"),
    ("Patents", "/company/patents", "patents.jpg"),
], cols=2, img_w=4.6)

# Industry children (11)
doc.add_page_break()
section_heading(doc, "Industry Use Cases  \u2014  /industries/[slug]", space_before=4)
add_grid(doc, [
    ("Manufacturing", "/industries/manufacturing", "ind-manufacturing.jpg"),
    ("Oil & Gas", "/industries/oil-gas", "ind-oil-gas.jpg"),
    ("Utilities", "/industries/utilities", "ind-utilities.jpg"),
    ("Financial Services", "/industries/financial-services", "ind-financial-services.jpg"),
    ("Government and Defense", "/industries/government-and-defense", "ind-government-and-defense.jpg"),
    ("Maritime", "/industries/maritime", "ind-maritime.jpg"),
    ("Life Sciences and Healthcare", "/industries/life-sciences-and-healthcare", "ind-life-sciences-and-healthcare.jpg"),
    ("Retail", "/industries/retail", "ind-retail.jpg"),
    ("Telecommunications", "/industries/telecommunications", "ind-telecommunications.jpg"),
    ("Transportation", "/industries/transportation", "ind-transportation.jpg"),
    ("Information Technology & Agentic AI", "/industries/information-technology-agentic-ai", "ind-information-technology-agentic-ai.jpg"),
], cols=3, img_w=3.05)

out = "AMERA-Visual-Site-Map.docx"
doc.save(out)
print("saved", out)
