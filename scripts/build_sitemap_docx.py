from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# --- Site map data ---------------------------------------------------------
HOME = ("Home", "/")

COLUMNS = [
    ("Platform", "/products", [
        ("AmeraKey", "/products/amerakey"),
        ("AmeraSecrets", "/products/amerasecrets"),
    ]),
    ("Industry Use Cases", "/industries", [
        ("Manufacturing", "/industries/manufacturing"),
        ("Oil & Gas", "/industries/oil-gas"),
        ("Utilities", "/industries/utilities"),
        ("Financial Services", "/industries/financial-services"),
        ("Government and Defense", "/industries/government-and-defense"),
        ("Maritime", "/industries/maritime"),
        ("Life Sciences and Healthcare", "/industries/life-sciences-and-healthcare"),
        ("Retail", "/industries/retail"),
        ("Telecommunications", "/industries/telecommunications"),
        ("Transportation", "/industries/transportation"),
        ("Information Technology & Agentic AI", "/industries/information-technology-agentic-ai"),
    ]),
    ("Resources", "/resources", []),
    ("News", "/news", []),
    ("Company", "/company", [
        ("Vision & Mission", "/company/vision-and-mission"),
        ("Executive Team", "/company/executive-team"),
        ("Board of Advisors", "/company/board-of-advisors"),
        ("Patents", "/company/patents"),
    ]),
    ("Contact Us", "/contact", []),
]

# Colors
HOME_FILL = "2F6F9F"        # deep blue
TOP_FILL = "4D9FD6"         # primary blue
CHILD_FILL = "DCEEF9"       # light blue
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1F, 0x2A, 0x33)


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_border(cell, color="FFFFFF", sz="12"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def vertical_center(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), "center")
    tcPr.append(va)


def style_box_cell(cell, title, path, fill, title_color, path_color,
                   title_size=10, path_size=8, border="FFFFFF"):
    set_cell_bg(cell, fill)
    set_cell_border(cell, color=border)
    vertical_center(cell)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(title_size)
    r.font.color.rgb = title_color
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run(path)
    r2.font.size = Pt(path_size)
    r2.font.color.rgb = path_color


# --- Build document --------------------------------------------------------
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
# swap dimensions for landscape
sec.page_width, sec.page_height = Inches(11), Inches(8.5)
sec.left_margin = sec.right_margin = Inches(0.4)
sec.top_margin = sec.bottom_margin = Inches(0.4)

# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = title.add_run("AMERA Website \u2014 Site Map")
tr.bold = True
tr.font.size = Pt(18)
tr.font.color.rgb = DARK
title.paragraph_format.space_after = Pt(6)

# Home box (single-cell centered table, ~3in wide)
home_tbl = doc.add_table(rows=1, cols=1)
home_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
home_tbl.autofit = False
hc = home_tbl.cell(0, 0)
hc.width = Inches(3.0)
style_box_cell(hc, HOME[0], HOME[1], HOME_FILL, WHITE, RGBColor(0xE8, 0xF2, 0xFB),
               title_size=13, path_size=9)

# connector arrow
conn = doc.add_paragraph()
conn.alignment = WD_ALIGN_PARAGRAPH.CENTER
cr = conn.add_run("\u2193")
cr.font.size = Pt(12)
cr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
conn.paragraph_format.space_before = Pt(0)
conn.paragraph_format.space_after = Pt(2)

# Main table: one row of top-level boxes, then child rows
n_cols = len(COLUMNS)
max_children = max(len(c[2]) for c in COLUMNS)
n_rows = 1 + max_children

tbl = doc.add_table(rows=n_rows, cols=n_cols)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.autofit = False
col_w = Inches(10.2 / n_cols)

# header row (top-level pages)
for ci, (name, path, children) in enumerate(COLUMNS):
    cell = tbl.cell(0, ci)
    cell.width = col_w
    style_box_cell(cell, name, path, TOP_FILL, WHITE,
                   RGBColor(0xE8, 0xF2, 0xFB), title_size=10, path_size=8)

# child rows
for ci, (name, path, children) in enumerate(COLUMNS):
    for ri in range(max_children):
        cell = tbl.cell(1 + ri, ci)
        cell.width = col_w
        if ri < len(children):
            cname, cpath = children[ri]
            style_box_cell(cell, cname, cpath, CHILD_FILL, DARK, DARK,
                           title_size=8.5, path_size=7, border="FFFFFF")
        else:
            # empty filler cell: white, no visible border
            set_cell_bg(cell, "FFFFFF")
            set_cell_border(cell, color="FFFFFF")

# Footnote
note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(8)
nr = note.add_run("All 11 industry pages are served by the dynamic /industries/[slug] template.")
nr.italic = True
nr.font.size = Pt(8)
nr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

out = "AMERA-Site-Map.docx"
doc.save(out)
print("saved", out)
